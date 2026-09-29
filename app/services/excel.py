from pathlib import Path
from datetime import datetime
from typing import Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import CATALOGUE_FINAL_DIR

EXCEL_PATH = CATALOGUE_FINAL_DIR / "catalogue_global.xlsx"

def init_excel_workbook() -> openpyxl.Workbook:
    """Creates a new openpyxl Workbook initialized with styled headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalogue Produits"
    
    headers = ["Catégorie", "Nom Fichier", "Prix", "Chemin d'Accès", "Date Traitement"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='medium', color='111827')
    )
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 25
    return wb

def ensure_excel_registry_exists() -> Path:
    """
    Ensures that the catalogue_global.xlsx file exists with proper headers.
    Returns the Path to the Excel file.
    """
    CATALOGUE_FINAL_DIR.mkdir(parents=True, exist_ok=True)
    if not EXCEL_PATH.exists():
        wb = init_excel_workbook()
        # Default column widths
        for col_letter in ["A", "B", "C", "D", "E"]:
            wb.active.column_dimensions[col_letter].width = 20
        wb.save(EXCEL_PATH)
    return EXCEL_PATH

def append_to_excel_registry(category_name: str, product_filename: str, file_path: str, price: str = "") -> str:
    """
    Appends a new processed product record to catalogue_final/catalogue_global.xlsx.
    Creates the workbook with styled headers if it does not exist yet.
    """
    CATALOGUE_FINAL_DIR.mkdir(parents=True, exist_ok=True)
    
    if EXCEL_PATH.exists():
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active
    else:
        wb = init_excel_workbook()
        ws = wb.active

    # Append Data Row
    timestamp_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    row_data = [category_name, product_filename, price or "-", file_path, timestamp_str]
    ws.append(row_data)
    
    row_idx = ws.max_row
    data_font = Font(name="Calibri", size=10)
    data_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    for col_num in range(1, len(row_data) + 1):
        cell = ws.cell(row=row_idx, column=col_num)
        cell.font = data_font
        cell.border = data_border
        if col_num in (1, 3, 5):
            cell.alignment = Alignment(horizontal="center")
        else:
            cell.alignment = Alignment(horizontal="left")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(EXCEL_PATH)
    return str(EXCEL_PATH)
