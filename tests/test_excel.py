import openpyxl
from pathlib import Path
from app.services.excel import ensure_excel_registry_exists, append_to_excel_registry

def test_ensure_excel_registry_exists(isolated_env):
    excel_path = isolated_env["excel_path"]
    assert excel_path.exists()

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    assert ws.title == "Catalogue Produits"
    headers = [cell.value for cell in ws[1]]
    assert headers == ["Catégorie", "Nom Fichier", "Prix", "Chemin d'Accès", "Date Traitement"]

def test_append_to_excel_registry(isolated_env):
    excel_path = isolated_env["excel_path"]
    
    append_to_excel_registry(
        category_name="Tartes",
        product_filename="tarte_citron_1.png",
        file_path="catalogue_final/tartes/tarte_citron_1.png",
        price="12 000 FCFA"
    )

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    assert ws.max_row == 2
    row2 = [cell.value for cell in ws[2]]
    assert row2[0] == "Tartes"
    assert row2[1] == "tarte_citron_1.png"
    assert row2[2] == "12 000 FCFA"
    assert "catalogue_final/tartes/tarte_citron_1.png" in row2[3]
